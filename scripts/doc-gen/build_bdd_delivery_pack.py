from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT=Path('/Users/dgsmacbook/Documents/Codex/2026-08-12/referenced-chatgpt-conversation-this-is-an/outputs/BDD_Final_Deliverables_and_Submission_Pack.docx')
OUT.parent.mkdir(parents=True,exist_ok=True)
NAVY='102A43'; BLUE='146C94'; TEAL='0E9F9A'; GREY='52616B'
def shade(c,x):
 p=c._tc.get_or_add_tcPr(); e=OxmlElement('w:shd'); e.set(qn('w:fill'),x); p.append(e)
def border(c):
 p=c._tc.get_or_add_tcPr(); b=OxmlElement('w:tcBorders')
 for n in ('top','left','bottom','right'):
  e=OxmlElement('w:'+n); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),'D0D7DE'); b.append(e)
 p.append(b)
def cell(c,t,b=False,col=None):
 c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(2); r=p.add_run(str(t)); r.bold=b; r.font.size=Pt(8.2)
 if col:r.font.color.rgb=RGBColor.from_string(col)
 c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP;border(c)
def table(d,h,rows):
 t=d.add_table(rows=1,cols=len(h));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,x in enumerate(h):shade(t.rows[0].cells[i],NAVY);cell(t.rows[0].cells[i],x,True,'FFFFFF')
 for ri,row in enumerate(rows):
  cs=t.add_row().cells
  for i,x in enumerate(row):
   if ri%2:shade(cs[i],'F6F8FA')
   cell(cs[i],x)
 return t
def p(d,t='',style=None):return d.add_paragraph(t,style)
def h(d,t,l=1):return d.add_heading(t,l)
def b(d,t):d.add_paragraph(t,style='List Bullet')
def code(d,t):
 q=d.add_paragraph();q.paragraph_format.left_indent=Inches(.2);q.paragraph_format.right_indent=Inches(.2);q.paragraph_format.space_before=Pt(4);q.paragraph_format.space_after=Pt(6);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F1F5F9');q._p.get_or_add_pPr().append(sh);r=q.add_run(t);r.font.name='Menlo';r.font.size=Pt(8)
d=Document();s=d.sections[0];s.top_margin=Inches(.7);s.bottom_margin=Inches(.7);s.left_margin=Inches(.75);s.right_margin=Inches(.75)
d.styles['Normal'].font.name='Aptos';d.styles['Normal'].font.size=Pt(10)
for n,z,c in [('Title',29,NAVY),('Heading 1',18,NAVY),('Heading 2',13,BLUE),('Heading 3',11,TEAL)]:
 st=d.styles[n];st.font.name='Aptos Display';st.font.size=Pt(z);st.font.bold=True;st.font.color.rgb=RGBColor.from_string(c);st.paragraph_format.space_before=Pt(15);st.paragraph_format.space_after=Pt(6)
head=s.header.paragraphs[0];head.text='BHARAT DATA DETECTIVE  |  FINAL DELIVERABLES';head.runs[0].font.size=Pt(8);head.runs[0].font.color.rgb=RGBColor.from_string(GREY)
foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER;foot.add_run('BDD • Final Deliverables & Submission Pack  |  ');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
q=d.add_paragraph();q.paragraph_format.space_before=Pt(56);q.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=q.add_run('BDD FINAL DELIVERABLES');r.font.name='Aptos Display';r.font.size=Pt(28);r.bold=True;r.font.color.rgb=RGBColor.from_string(NAVY)
q=d.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=q.add_run('Build, Evidence, Evaluation & Competition Submission Pack');r.font.size=Pt(18);r.font.color.rgb=RGBColor.from_string(BLUE)
q=d.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_before=Pt(18);q.add_run('A practical checklist for turning Bharat Data Detective into a credible, demo-ready submission.').italic=True
table(d,['Document control','Value'],[['Purpose','Defines exactly what the team must produce and how to prove it is real.'],['Use with','BDD Project Specification v1.0 (the technical blueprint).'],['Delivery principle','Every claim in a deck/demo must be traceable to a working artifact, evidence bundle or explicitly labeled prototype limitation.'],['Scope','V1: one programme domain, batch workflow, 2–3 benchmark candidate cases, public/approved data only.']])
d.add_page_break()
h(d,'1. Delivery Outcome')
p(d,'The finished BDD submission is a reproducible forensic-data product, not only a dashboard. A reviewer should be able to open the app, inspect an approved public source, follow a potential inconsistency back through transformations to evidence, see what remains uncertain, and review an honest blinded-benchmark result.')
table(d,['Delivery track','What “done” looks like'],[
['A. Product','A locally runnable web application and API that ingest/profile approved fixtures, generate findings, display evidence and answer scoped questions with citations.'],
['B. Evidence','An immutable provenance pack for every demo source plus a restricted, independently cited CAG benchmark registry.'],
['C. Evaluation','A documented blinded run and adjudication outcome for each selected case; metrics only if definitions and denominators support them.'],
['D. Competition','A concise narrative, recorded/live demo, slides and one-page summary that show BDD’s four signature innovations responsibly.']])
h(d,'2. Master Deliverables Register')
table(d,['#','Deliverable','Owner','Required evidence / completion test'],[
['D1','Working prototype repository','Engineering lead','Clean clone + documented command launches stack; CI green; no secrets committed.'],
['D2','Public demo corpus and source register','Data lead','Each artifact has URL, publisher, retrieval time, licence/terms note, hash and local version.'],
['D3','CAG benchmark registry (restricted)','Evaluation lead','Each case cites report/page/paragraph; labels isolated from investigator data and retrieval.'],
['D4','Forensic case files (2–3)','Forensics lead','Each case has scope, finding(s), evidence graph/export, competing explanations, reviewer disposition and reproducibility recipe.'],
['D5','Evaluation report','Evaluation lead','Predefined matching policy, coverage, case outcomes, limitations and metrics only where valid.'],
['D6','Technical documentation','Engineering/data leads','Architecture, setup, APIs, schemas, ADRs, security/provenance and operations notes.'],
['D7','Competition deck + one-pager','Product/presentation lead','Uses measured claims only; screenshots and demo evidence match product.'],
['D8','Demo video or live-demo runbook','Demo lead','5–7 minute path with offline fallback and timing rehearsed.'],
['D9','Submission archive','Project lead','Read-only tagged release, checksums, licence list, submission links and final checklist signed.']])
h(d,'3. Artifact-by-Artifact Requirements')
h(d,'3.1 D1 — Working Prototype',2)
table(d,['Component','Minimum content'],[
['Web dashboard','Dataset list, dataset page, findings queue, case view, evidence panel, comparability/drift signal and Ask Detective.'],
['API','Typed endpoints for source/artifact ingestion, profiling, mappings, rules/runs, findings, evidence and scoped Ask Detective.'],
['Pipelines','Ingestion → immutable manifest → profiling → normalization → deterministic checks → evidence graph → case file.'],
['Storage','Postgres/PostGIS metadata; object storage/raw filesystem; Qdrant for retrieved evidence; NetworkX graph projection.'],
['Safety controls','RBAC/role stub, benchmark isolation, citation validation, prompt-injection handling and no unsupported conclusion state.'],
['Reproducibility','Docker Compose, pinned dependencies, migrations, seed fixtures and `make demo`/equivalent command.']])
h(d,'3.2 D2 — Public Demo Corpus & Source Register',2)
p(d,'Use official ministry/department release as primary where available, then data.gov.in and AIKosh for discovery/alternate distribution. Preserve a raw immutable copy and never overwrite an older version.')
code(d,'source_id,artifact_id,official_title,publisher,source_url,resource_url,retrieved_at,release_date,\nsha256,media_type,licence_or_terms,programme,geo_coverage,time_coverage,access_status,notes')
b(d,'Minimum corpus: at least two related structured sources or two releases of one source, plus their human-readable metadata/data dictionary where available.')
b(d,'Every demo screen must show a source locator and retrieval metadata; every transformation must retain raw-to-derived lineage.')
h(d,'3.3 D3 — Restricted CAG Benchmark Registry',2)
p(d,'This registry is an evaluation control. Do not place it in the public application seed data, vector database, agent prompt files or accessible demo folder. Keep it in a restricted path/repository or a separate protected database schema.')
code(d,'benchmark_case_id,audit_report_title,report_url,audit_period,page_or_paragraph,observation_summary,\nlinked_source_ids,label_visibility,detectability_status,matching_policy_version,adjudicator,decision_date')
table(d,['Control','Required practice'],[['Blindness','Investigator receives source map and permitted inputs, not the audit conclusion/count/label.'],['Linkage','Record why each public source could plausibly expose the audited condition.'],['Adjudication','Reveal label only after run; record detected / partly detected / not detected / not detectable and rationale.'],['Integrity','Record immutable benchmark-version hash and access log.']])
h(d,'3.4 D4 — Forensic Case File',2)
table(d,['Case-file section','Required content'],[
['Question and scope','What is being tested; programme, sources, dates, geography, filters and exclusions.'],['Observed facts','Only directly reproducible facts: e.g., counts/deltas/null rates with query/rule references.'],['Definitions and comparability','Units, denominators, reporting cut-offs, geography/boundary version and drift assessment.'],['Finding','Safe claim wording, severity, confidence, alternatives and why it needs review.'],['Evidence','Artifact IDs/hashes, source locators, row/aggregate references, graph path and citations.'],['Reproduction','Pinned configuration, SQL/rule version, model/prompt version if applicable and run ID.'],['Review','Human decision, rationale, reviewer and date.'],['Benchmark outcome','Restricted until blind run is complete; then state outcome and limits.']])
h(d,'3.5 D5 — Evaluation Report',2)
p(d,'The report must state the candidate-case universe, accessible-data coverage, what was hidden, evaluation matching rules, and all exclusions. A small number of cases may make F1 uninformative; in that situation report the case table and do not oversell a metric.')
table(d,['Metric / result','Required definition'],[['Case coverage','Cases with enough permissible public data ÷ registered candidate cases.'],['Detectability rate','Cases adjudicated detectable ÷ cases with accessible data.'],['Precision / recall / F1','Only after pre-registering TP/FP/FN matching criteria and reviewer adjudication.'],['Evidence completeness','Findings with artifact + locator + reproduction recipe ÷ displayed findings.'],['Citation correctness','Sampled answer claims whose citation directly supports the claim ÷ sampled cited claims.'],['Time to investigation','Wall-clock time from selected case to reviewer-ready case file; state hardware/config.']])
h(d,'3.6 D6 — Technical Documentation',2)
table(d,['Document','Minimum contents'],[['README','What BDD is/is not, local setup, demo command, architecture image/flow and troubleshooting.'],['Architecture & API guide','Service boundaries, OpenAPI location, data flow, engine contracts, async jobs and failure handling.'],['Data governance guide','Source approval, licence/terms, raw retention, PII policy, benchmark access, citations and deletion/change policy.'],['Model/RAG card','Candidate/model selection, license, model hash/config, benchmarks, retrieval policy, known limitations and safety gates.'],['ADR folder','Why source policy, scoring weights, graph store, model and V1 scope were chosen.'],['Runbooks','Ingest source, add rule, investigate finding, reveal benchmark, roll back/rebuild index and demo offline.']])
h(d,'4. Recommended Submission Folder')
code(d,'bdd-submission/\n  README.md\n  docker-compose.yml\n  .env.example\n  apps/  services/  pipelines/  packages/\n  docs/\n    architecture.md  api.md  governance.md  model-card.md  runbooks/  adr/\n  data/\n    public-fixtures/  source-register.csv  manifests/\n    benchmark-restricted/  # exclude from public release and retrieval\n  cases/\n    CASE-001/ case-file.md evidence-bundle.json reproduction.json\n  evaluation/\n    protocol.md adjudications.csv report.pdf\n  competition/\n    deck.pdf one-pager.pdf demo-script.md screenshots/ video-link.txt\n  infra/  tests/  LICENSE  NOTICE')
h(d,'5. Team Operating Cadence')
table(d,['Cadence','Agenda / output'],[['Daily 15 min','Blockers, source/data decisions, demo-risk check.'],['Source review','Approve or reject new sources; verify licence, provenance and benchmark isolation.'],['Forensics review','Review findings for comparability, evidence completeness, safe wording and human disposition.'],['Weekly demo rehearsal','Run from clean environment; time it; record gaps in a visible issue list.'],['Release gate','Confirm clean build, cited claims, evaluation integrity, security scan and submission checklist.']])
h(d,'6. Final Acceptance Checklist')
checks=['The project runs from a clean, documented environment with a single demo command.','At least one end-to-end investigation is shown from raw artifact to cited case file.','All demo sources are public/approved and have manifests, hashes and source locators.','The UI distinguishes potential inconsistencies from verified errors and exposes confidence/severity separately.','A cross-source or temporal comparison passes a definition/unit/time/geography comparability gate.','The evidence graph shows a trace from finding to raw evidence and transformation/rule run.','False consensus view explains common origin versus independent corroboration.','Semantic drift view compares at least two versions/definitions or explicitly states unavailable evidence.','Ask Detective is scoped, citation-first and refuses/qualifies unsupported questions.','CAG benchmark content was isolated during investigation; adjudication records exist.','No performance result is claimed unless a measured report with denominator and methodology exists.','Secrets are absent from repository; `.env.example` contains placeholders only.','Deck/demo uses the exact same evidence and wording as the product.','A backup offline demo path (recorded video or fixtures) is tested.','Release tag/archive, licences and final submission links are captured.']
for i,x in enumerate(checks,1):p(d,f'☐ {i}. {x}')
h(d,'7. Demo Script: 6 Minutes')
table(d,['Time','Action','Message'],[['0:00–0:40','Problem and promise','Public data needs evidence trust, not just more dashboards or an LLM answer.'],['0:40–1:20','Ingest source','Official artifact is captured with hash, source metadata and raw preservation.'],['1:20–2:10','Profile and understand','Show schema, missingness and definition cards; point out comparable/non-comparable fields.'],['2:10–3:20','Investigate','Run deterministic comparison/anomaly; reveal a potential inconsistency and alternative explanations.'],['3:20–4:10','Trace evidence','Follow the lineage graph from finding to source/aggregate and rule.'],['4:10–4:50','Ask Detective','Ask a scoped question; show cited answer, uncertainty and next action.'],['4:50–5:30','Signature innovations','Drift, lineage, false consensus, cross-source forensics.'],['5:30–6:00','Honest evaluation','Explain blinded CAG benchmark and measured status/limitations.']])
h(d,'8. Competition Claims: Allowed vs Not Allowed')
table(d,['Use in submission','Avoid unless independently established'],[['“BDD flags evidence-backed potential inconsistencies requiring verification.”','“BDD proves government data is wrong.”'],['“The demo traces every displayed finding to an artifact, locator and rule run.”','“AI detects fraud/eligibility errors.”'],['“CAG reports provide an independent historical benchmark protocol.”','“BDD has achieved X% accuracy” without results and method.'],['“Multiple sources may share one origin; BDD measures evidence diversity.”','“Three websites mean three independent confirmations.”'],['“Current case result: detected / not detectable / under adjudication.”','“Audit observation was reproduced” before blinded adjudication.']])
h(d,'9. Handoff Package')
p(d,'When handing BDD to another student, mentor or coding agent, provide the repository URL/release archive; the current specification; this delivery pack; source register; restricted benchmark access instructions; current case-file index; test status; known limitations; and the next three prioritized tasks. The handoff should make it possible to continue work without re-discovering data provenance or silently weakening safety controls.')
h(d,'10. Final Submission Sign-off')
table(d,['Role','Name','Date','Sign-off statement'],[['Project lead','','','Scope, deliverables and submission artifacts checked.'],['Data/provenance lead','','','Sources, hashes, licences and citations checked.'],['Evaluation lead','','','Benchmark isolation and reporting methodology checked.'],['Engineering lead','','','Build, tests, security basics and reproducibility checked.'],['Presentation lead','','','Deck/demo claims match documented evidence.']])
d.core_properties.title='BDD Final Deliverables and Submission Pack';d.core_properties.author='BDD Project Team';d.save(OUT);print(OUT)
