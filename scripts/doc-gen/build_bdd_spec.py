from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parents[2] / "docs" / "specification" / "Bharat_Data_Detective_BDD_Project_Specification.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY='102A43'; BLUE='146C94'; TEAL='0E9F9A'; GOLD='DFAE35'; PALE='EAF3F7'; GREY='52616B'; RED='B42318'; GREEN='027A48'

def shade(cell, color):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),color); tcPr.append(shd)
def borders(cell, color='D0D7DE'):
    tcPr=cell._tc.get_or_add_tcPr(); b=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
        x=OxmlElement('w:'+edge); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'6'); x.set(qn('w:color'),color); b.append(x)
    tcPr.append(b)
def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); e=OxmlElement('w:tblHeader'); e.set(qn('w:val'),'true'); trPr.append(e)
def set_cell_text(cell, text, bold=False, color=None, size=8.5):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2); r=p.add_run(str(text)); r.bold=bold; r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP; borders(cell)
def add_table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    h=t.rows[0]; set_repeat_table_header(h)
    for i,x in enumerate(headers): shade(h.cells[i],NAVY); set_cell_text(h.cells[i],x,True,'FFFFFF',8.3)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,x in enumerate(row):
            if ri%2==1: shade(cells[i],'F6F8FA')
            set_cell_text(cells[i],x,False,None,8.1)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def bullet(doc, text, level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.add_run(text); return p
def para(doc, text='', style=None, boldlead=None):
    p=doc.add_paragraph(style=style)
    if boldlead and text.startswith(boldlead):
        p.add_run(boldlead).bold=True; p.add_run(text[len(boldlead):])
    else:p.add_run(text)
    return p
def heading(doc, text, level=1): return doc.add_heading(text,level)
def page_break(doc): doc.add_page_break()
def codeblock(doc,text):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.22); p.paragraph_format.right_indent=Inches(.22); p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(7)
    shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'F1F5F9'); p._p.get_or_add_pPr().append(shd); r=p.add_run(text); r.font.name='Menlo'; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(NAVY)

d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
styles=d.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10); styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Aptos')
for n,size,col in [('Title',30,NAVY),('Heading 1',18,NAVY),('Heading 2',13,BLUE),('Heading 3',11,TEAL)]:
    s=styles[n]; s.font.name='Aptos Display' if n!='Heading 3' else 'Aptos'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(col); s.font.bold=True
    s.paragraph_format.space_before=Pt(16 if n!='Heading 3' else 10); s.paragraph_format.space_after=Pt(6)
if 'BDD Callout' not in styles:
    s=styles.add_style('BDD Callout',WD_STYLE_TYPE.PARAGRAPH); s.font.name='Aptos'; s.font.size=Pt(10); s.font.color.rgb=RGBColor.from_string(NAVY); s.paragraph_format.left_indent=Inches(.22); s.paragraph_format.right_indent=Inches(.22); s.paragraph_format.space_before=Pt(7); s.paragraph_format.space_after=Pt(7)
# header/footer
header=sec.header.paragraphs[0]; header.text='BHARAT DATA DETECTIVE  |  PROJECT SPECIFICATION'; header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor.from_string(GREY)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run('BDD • AI Forensic & Evidence Trust Layer for Indian Public Data  |  '); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

# Cover
p=d.add_paragraph(); p.paragraph_format.space_before=Pt(60); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('BHARAT DATA DETECTIVE'); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(31); r.font.color.rgb=RGBColor.from_string(NAVY)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('BDD'); r.bold=True; r.font.size=Pt(17); r.font.color.rgb=RGBColor.from_string(TEAL)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('AI Forensic & Evidence Trust Layer for Indian Public Data'); r.font.size=Pt(19); r.font.color.rgb=RGBColor.from_string(BLUE)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(22); p.add_run('Implementation-ready project specification and competition narrative').italic=True
add_table(d,['Document control','Value'],[['Version','1.0 — implementation baseline'],['Audience','Student teams, coding agents, reviewers, mentors and hackathon juries'],['Status','Build specification; performance claims intentionally excluded until measured'],['Core principle','BDD identifies evidence-backed inconsistencies for review. It does not declare a source or person “wrong.”'],['Signature innovations','Semantic Drift Detection • Evidence Lineage Graph • False Consensus Detection • Cross-source AI Forensics']],[1.8,5.8])
para(d,'Prepared for an India-first public-data trust system. This document assumes public, lawfully obtained datasets and a human-in-the-loop forensic workflow.', 'BDD Callout')
page_break(d)

heading(d,'Table of Contents',1); p=d.add_paragraph(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'TOC \\o "1-3" \\h \\z \\u'); p._p.append(fld); para(d,'Update fields in Word (right-click → Update Field) to refresh page numbers after edits.', 'BDD Callout'); page_break(d)

heading(d,'1. Executive Summary')
para(d,'Bharat Data Detective (BDD) is an AI-assisted forensic and evidence-trust layer for Indian public data. It ingests public datasets and accompanying documents, profiles their structure and quality, resolves what fields mean, compares evidence across sources and time, and produces reviewable case files. Its output is not a truth verdict: it is an evidence-backed finding with source lineage, reproducible checks, uncertainty, and a recommended verification action.')
para(d,'The V1 product is deliberately narrow: reproduce the detectability of two to three historically documented public-data-quality cases using public data available to BDD, while keeping the relevant Comptroller and Auditor General of India (CAG) observations hidden from the investigation workflow until evaluation. This makes the project more than a synthetic-error demo and prevents unsupported accuracy claims.')
add_table(d,['What BDD does','What BDD does not do'],[
['Flags potential inconsistencies, anomalies, definition changes and contradictions','Accuse beneficiaries, officials, agencies or datasets of fraud or falsity'],
['Preserves raw artifacts, hashes, transformations, queries and citations','Replace statutory audits, eligibility decisions or domain verification'],
['Ranks records/cells/aggregates for human review','Infer sensitive personal attributes or bypass access controls'],
['Explains why multiple sources may be falsely agreeing','Treat similarity, popularity or an LLM response as proof']])
heading(d,'2. Problem Evidence and Motivation')
para(d,'Indian public data is distributed across portals, department sites, PDFs, dashboards, spreadsheets and APIs. The same programme can appear in several copies with different update dates, granularities, definitions, units and geographic labels. Known audit observations show that administrative-data issues can have material consequences. BDD’s motivation is therefore practical: help analysts move from “a number looks suspicious” to a reproducible evidence trail that a domain owner can verify.')
para(d,'CAG reports are used as independent historical evidence, not as training labels or content injected into the investigator prompt. The initial corpus must contain only audit observations that can be linked to a public raw-data or public-document trail; inaccessible source data is still recorded as a coverage gap, not silently excluded.')
heading(d,'3. Vision, Principles and Non-Goals')
add_table(d,['Vision / principle','Operational consequence'],[
['Trust should be inspectable','Every finding links to artifact, version, row/aggregate, rule, model and timestamp.'],['Evidence before narrative','LLMs summarize bounded evidence; deterministic checks establish numeric facts.'],['India-first semantics','Geographic hierarchies, local nomenclature, schemes, fiscal years and multilingual fields are first-class.'],['Reviewability','A reviewer can rerun a case from pinned inputs and see why a score was assigned.'],['Safe uncertainty','Use “potential inconsistency”, “requires verification”, and calibrated confidence.']])
bullet(d,'Non-goal: build a universal fact checker or a replacement for CAG/departmental audit processes.')
bullet(d,'Non-goal: make causal or legal findings from public data alone.')
bullet(d,'Non-goal: begin V1 with unrestricted web crawling, live operational data, or personally identifying records.')

heading(d,'4. Users and Primary Use Cases')
add_table(d,['User','Job to be done','BDD outcome'],[
['Data analyst / student','Assess whether a public dataset is safe and intelligible for analysis or RAG','Profile, field definitions, fitness scores, caveats and citations'],['Programme researcher','Investigate a discrepancy across releases or sources','Case file with comparable measures, lineage, reconciliation and review steps'],['Data publisher / steward','Find quality regressions before publication','Validation report, changed-schema alert and remediation queue'],['Hackathon judge / mentor','Assess technical depth and responsible AI','Live, reproducible forensic investigation with hidden-benchmark evaluation'],['Downstream RAG builder','Decide whether a corpus can be retrieved safely','RAG Readiness Score, chunk-level provenance and exclusions']])

heading(d,'5. Exact V1 Scope')
para(d,'V1 supports batch ingestion of CSV/XLSX/JSON/Parquet and text-extracted PDF/HTML metadata, one pilot programme domain, 2–3 public-data cases, India administrative geographies through district (where source supports it), English plus transliterated/multilingual field aliases, and a human-reviewed investigation queue. It produces profiles, mappings, quality checks, cross-source comparisons, a lineage graph, case files, and a dashboard.')
add_table(d,['In V1','Deferred after V1'],[
['Offline/batch sources, signed manifests, manual source approval','Continuous crawling, auto-remediation or write-back to publisher systems'],
['Rule + statistical + bounded LLM investigation','Autonomous conclusions, agentic external actions or high-stakes decisions'],
['NetworkX graph projection; Neo4j-ready interface','Mandatory graph database at prototype scale'],
['One controlled model benchmark and one approved model','Fine-tuning or model-specific production lock-in']])

heading(d,'6. India-First Dataset and Benchmark Strategy')
heading(d,'6.1 Source hierarchy')
add_table(d,['Tier','Source class','Use','Required provenance'],[
['A','Original ministry/department source, official dashboard/API/public release','Primary artifact and authoritative release context','URL, publisher, retrieval time, hash, licence/terms, release/update date'],['B','data.gov.in catalogue/API','Discovery and alternate distribution; preserve resource linkage','Catalogue ID, resource URL, publisher, API parameters, retrieval timestamp'],['C','AIKosh catalogue / approved public research dataset','Discovery, metadata and permissible dataset access','Catalogue reference, owner, licence, exact asset URL'],['D','CAG audit report','Hidden independent benchmark / audit observation source','Report title, audit period, page/paragraph, URL, extraction hash'],['E','News/blog/secondary reporting','Context only; never sole evidence of a finding','Publisher/date/URL and link to primary source where available']])
para(d,'Source precedence is not truth precedence. A Tier A source may still contain an inconsistency; precedence only governs how BDD describes authority and reconciliation.')
heading(d,'6.2 Benchmark-corpus methodology')
add_table(d,['Step','Method','Artifact'],[
['1. Register','Create a case registry from CAG observations; record exact citation and audit period.','case_registry.csv (restricted benchmark labels)'],['2. Map','Locate candidate public datasets/releases and document the linkage rationale.','source_map.json'],['3. Freeze','Download immutable copies; calculate SHA-256; store raw and normalized forms separately.','manifest.json + raw artifacts'],['4. Blind','Withhold CAG conclusion text, outcome label and exact affected-record count from investigator inputs.','benchmark split + access policy'],['5. Investigate','Run BDD only on allowed raw/metadata inputs; reviewers log hypotheses before label reveal.','case run bundle'],['6. Adjudicate','Reveal benchmark; classify detected / partially detected / not detectable / not reproducible.','adjudication form'],['7. Measure','Calculate case-level and finding-level metrics with confidence intervals where sample size allows.','evaluation report']])
para(d,'A “not detectable” outcome is valuable: it identifies a data-access, observability, or semantic limitation rather than a model failure. Never convert an audit narrative into synthetic labels without marking it as an illustrative derived test.')

heading(d,'7. Functional Requirements')
add_table(d,['ID','Requirement','Acceptance signal'],[
['FR-01','Ingest approved public artifacts with immutable manifests.','Hash, source URL, timestamp, parser version and licence metadata are present.'],['FR-02','Profile each dataset and column.','Profile contains types, nulls, uniqueness, values, distribution, PII scan and quality observations.'],['FR-03','Resolve schema/semantic definitions with human approval.','Canonical mapping contains evidence, confidence, status and reviewer.'],['FR-04','Normalize temporal, geographic, units and identifiers without overwriting raw values.','Normalized view preserves raw-to-derived lineage.'],['FR-05','Generate deterministic anomalies and cross-source comparisons.','Every finding has rule/query/versioned inputs and reproducible result.'],['FR-06','Construct and query evidence lineage.','Case file graph can trace any assertion back to source spans/artifacts.'],['FR-07','Run bounded AI investigation and Ask Detective.','Answers use retrieved/pinned evidence, cite it, and state uncertainty.'],['FR-08','Score dataset fitness and RAG readiness.','Score breakdown and exclusions are inspectable.'],['FR-09','Evaluate against hidden CAG benchmark.','Metrics, denominators and case outcomes are reported without invented results.']])
heading(d,'8. Non-Functional Requirements')
add_table(d,['Area','Target / requirement'],[
['Reproducibility','Same artifact hashes + configuration + model version must reproduce deterministic outputs exactly.'],['Traceability','100% of displayed findings must contain at least one evidence link and one provenance record.'],['Security','Secrets outside repository; RBAC for benchmark labels; immutable append-only audit events.'],['Performance','V1 target: profile a 100 MB structured file on a developer machine with progress reporting; benchmark actual timings before publishing.'],['Availability','Local Docker composition works offline after images/models/artifacts are cached.'],['Accessibility','Keyboard-operable dashboard; clear severity labels not conveyed by colour alone.'],['Maintainability','Typed API contracts, migrations, tests, versioned rules and ADRs.']])

heading(d,'9. Engine Specifications: Inputs and Exact Outputs')
para(d,'All engines emit a common envelope: run_id, engine_id, engine_version, input_artifact_ids, config_hash, created_at, status, warnings, evidence_refs and output_payload. “Exact output” below means the required contract fields, not a claim that all fields will always have values.')
engines=[
('Data Ingestion Engine','Source registration, URL/file/API response, source metadata, retrieval policy','ArtifactManifest; raw object URI; SHA-256; format/parser details; extraction log; source citation; ingestion warnings'),
('Dataset Profiler','Raw/normalized table, sampling policy','DatasetProfile; ColumnProfile[]; quality_observations[]; PII/identifier hints; candidate keys; profile evidence'),
('Schema Intelligence','Column names/types/samples, metadata/docs, profile','SchemaMapping[]: raw_field, canonical_field, transform, confidence, evidence, approval_status; unresolved_fields[]'),
('Semantic Definition Engine','Field mappings, data dictionary, source text, approved ontology','DefinitionCard[]: concept, operational definition, scope, denominator, unit, evidence spans, ambiguity flags'),
('Geographic Intelligence','Location values, codes, source geography notes, boundary version','GeoResolution[]: raw_value, canonical_code/name, level, boundary_version, match_method, confidence, alternatives'),
('Temporal Intelligence','Date fields, reporting period text, release timestamps','TemporalResolution[]: event_period, publication_time, fiscal_year, granularity, timezone, comparability flags'),
('Statistical Anomaly Engine','Metric series/table, dimensions, approved rules/baseline','AnomalyFinding[]: metric, slice, observed, expected/baseline, method, score, severity, evidence query, caveats'),
('Cross-Source Contradiction Engine','Comparable DefinitionCards + normalized aggregates from ≥2 sources','ContradictionFinding[]: claim A/B, reconciliation status, delta, alignment tests, possible explanations, evidence'),
('Evidence / Data Lineage Graph','Artifacts, records/aggregates, transformations, checks, claims, citations','Graph nodes/edges; provenance paths; claim support/contradiction sets; exportable case subgraph'),
('False Consensus Detector','Claim cluster, source lineage, publishing/reuse links, timestamps','ConsensusAssessment: independence score, shared-origin cluster, evidence diversity, suspected copy chain, warning'),
('Semantic Drift Detector','Versioned DefinitionCards/schemas/release notes across time','DriftFinding[]: changed term/field, before/after definitions, impact scope, comparability decision, evidence'),
('AI Investigation Agent','Pinned finding(s), retrieved evidence, permitted tools, investigation plan','InvestigationMemo: hypotheses, tool calls, cited observations, unresolved questions, recommended review; no ungrounded verdict'),
('Evidence Engine','Finding/case, cited artifacts, policy rules','EvidenceBundle: citations, source excerpts/row pointers, hashes, reproducibility recipe, claim language status'),
('Data Fitness Score','Profile, semantic completeness, lineage, freshness and usability policy','DataFitnessScore: 0–100, dimension scores, blockers, rationale, policy version'),
('RAG Readiness Score','Fitness output, chunkability, provenance coverage, sensitive-data policy','RAGReadinessScore: 0–100, allowed/excluded fields, chunk plan, citation coverage, blockers'),
('Ask Detective Interface','Natural-language question, active dataset/case scope','Answer: direct response, cited evidence cards, confidence, assumptions, recommended next step, refusal if unsupported')]
add_table(d,['Engine','Required inputs','Required outputs'],engines,[1.35,2.55,3.9])

heading(d,'10. Signature Innovations')
heading(d,'10.1 Semantic Drift Detection')
para(d,'BDD versions definitions, field names, valid values, units, denominators and geographical boundaries. A numerical change is not treated as comparable until the system checks whether the underlying concept remained comparable. Drift is scored by semantic impact, not merely string difference.')
heading(d,'10.2 Evidence Lineage Graph')
para(d,'BDD models the path from a source artifact through extraction, normalization, aggregation and rule execution to a displayed claim. The graph is the auditability backbone: a user can traverse backwards from a dashboard card to raw evidence, or forward from a source version to every affected case.')
heading(d,'10.3 False Consensus Detection')
para(d,'Several pages repeating a number are not independent confirmation if they trace back to the same release, scrape, press note or copied spreadsheet. BDD clusters claims by lineage, timing, textual fingerprints and explicit citations, then reports diversity of evidence rather than raw source count.')
heading(d,'10.4 Cross-source AI Forensics')
para(d,'Cross-source comparison is a staged process: deterministic alignment of entity, time, geography, unit, denominator and release; numerical reconciliation; then LLM-assisted interpretation of bounded source excerpts. The LLM proposes hypotheses and asks for missing evidence—it does not calculate facts or decide which source is “true.”')

heading(d,'11. Canonical Data and Metadata Schemas')
heading(d,'11.1 Artifact manifest (JSON)')
codeblock(d,'{\n  "artifact_id": "art_01H...", "source_id": "src_ministry_x",\n  "source_url": "https://…", "retrieved_at": "ISO-8601",\n  "sha256": "…", "media_type": "text/csv", "byte_size": 0,\n  "publisher": "…", "license_or_terms": "…", "release_date": "YYYY-MM-DD",\n  "parser": {"name":"polars_csv","version":"…","config_hash":"…"},\n  "supersedes_artifact_id": null, "raw_uri": "object://raw/…"\n}')
heading(d,'11.2 Canonical observation record')
codeblock(d,'{\n  "observation_id":"obs_…", "artifact_id":"art_…", "raw_locator":{"row":42},\n  "dimensions":{"geo_code":"IN-…","period":"2025-26","scheme":"…"},\n  "measure":{"concept_id":"beneficiary_count","raw_value":"1,234",\n    "normalized_value":1234, "unit":"count", "denominator":null},\n  "transform_ids":["tr_…"], "quality_flags":["missing_identifier"]\n}')
heading(d,'11.3 Finding and claim')
codeblock(d,'{\n  "finding_id":"fin_…", "finding_type":"cross_source_contradiction",\n  "status":"open", "severity":"high", "confidence":"moderate",\n  "claim_text":"Potential inconsistency requiring verification: …",\n  "evidence_refs":["ev_…"], "repro_recipe_id":"rr_…",\n  "alternative_explanations":["different reporting cut-off"],\n  "human_review":{"required":true,"state":"pending"}\n}')
heading(d,'11.4 Evidence graph entities and relationships')
add_table(d,['Entity','Key relations'],[
['Source','PUBLISHES → Artifact; DESCRIBES → Definition'],['Artifact','SUPERSEDES / DERIVED_FROM → Artifact; CONTAINS → Observation; CITES → Citation'],['Transformation','CONSUMES → Artifact/Observation; PRODUCES → Dataset/Observation'],['DefinitionCard','DEFINES → Concept; VERSION_OF → DefinitionCard; SUPPORTED_BY → Evidence'],['Finding','OBSERVES → Observation/Aggregate; GENERATED_BY → RuleRun; SUPPORTED_BY / CONTRADICTED_BY → Evidence'],['Claim','ABOUT → Concept; ASSERTED_IN → Artifact; DEPENDS_ON → Claim; SAME_ORIGIN_AS → Claim'],['Case','CONTAINS → Finding; EVALUATED_AGAINST → BenchmarkCase; REVIEWED_BY → User'],['Evidence','POINTS_TO → artifact locator; HAS_HASH → Artifact; SUPPORTS / QUALIFIES / REFUTES → Claim']])

heading(d,'12. Architecture and End-to-End Flow')
codeblock(d,'Approved sources → Ingestion & immutable raw store → Parser/normalizer → Profiler\n     → Schema + semantic + geo + temporal intelligence → Canonical warehouse\n     → Rules/statistics/comparison engines → Findings + Evidence Graph → Case files/dashboard\n     ↘ Retrieval index + constrained LangGraph agent → Ask Detective (citations mandatory)\n     ↘ Hidden CAG benchmark (isolated) → evaluator → metrics & adjudication report')
add_table(d,['Layer','Recommended implementation','Responsibility'],[
['Experience','Next.js + TypeScript + Tailwind + shadcn/ui','Dashboard, case review, source viewer, Ask Detective'],['API/orchestration','FastAPI + Pydantic + LangGraph','Typed APIs, job state, agent guardrails, auth boundary'],['Data processing','Python + Polars/Pandas + DuckDB + Great Expectations','Batch transforms, profiling, validation, reproducible SQL'],['Core persistence','PostgreSQL + PostGIS; object storage','Metadata, cases, entities, spatial data, raw artifacts'],['Graph','NetworkX initially; Neo4j adapter later','Lineage traversal, claim/source dependency'],['Retrieval','Qdrant + embeddings + reranker','Evidence retrieval only; citation chunk IDs'],['Runtime','Docker Compose; CI; OpenTelemetry','Repeatable local setup, test automation, observability']])
heading(d,'12.1 Deterministic versus LLM boundary')
add_table(d,['Deterministic code must own','LLM may assist with','LLM must never do alone'],[
['Parsing, hashes, calculations, thresholds, joins, dedupe, unit conversion, metrics and scoring arithmetic','Field alias suggestions, definition extraction, hypothesis generation, evidence summarization, retrieval query reformulation','Fabricate citations, calculate an unverified result, override policy, label wrongdoing, access hidden benchmark labels'],
['Finding identifiers, lineage edges, rule versions and reproducibility recipes','Ask clarifying questions and explain alternatives','Silently change canonical mappings or publish findings without human-approved language']])

heading(d,'13. APIs, Services and Database Design')
add_table(d,['Service','Representative endpoints / events'],[
['Catalog & ingestion','POST /sources; POST /artifacts/ingest; GET /artifacts/{id}/manifest'],['Profile & semantics','POST /profiles; POST /mappings/propose; POST /definitions/approve'],['Forensics','POST /runs; GET /findings; POST /findings/{id}/review; POST /comparisons'],['Evidence graph','GET /cases/{id}/graph; GET /evidence/{id}; GET /reproduce/{id}'],['Assistant','POST /ask with scoped case/dataset ID; returns answer + evidence cards'],['Evaluation','POST /benchmarks/{id}/evaluate (restricted role); GET /evaluation-reports/{id}']])
add_table(d,['Database area','Core tables / collections'],[
['Catalog','sources, source_versions, artifacts, artifact_locations, licences, retrieval_events'],['Data/semantic','datasets, columns, profiles, concept_definitions, schema_mappings, geo_resolutions, temporal_resolutions'],['Forensics','rule_sets, rule_runs, finding_runs, findings, finding_evidence, review_decisions, cases'],['Lineage','graph_nodes, graph_edges, transformations, aggregation_specs, reproducibility_recipes'],['Retrieval','documents, chunks, embeddings, retrieval_logs, citation_spans'],['Evaluation/security','benchmark_cases, benchmark_labels (separate schema), adjudications, users, roles, audit_events']])

heading(d,'14. Confidence, Severity and Claim Taxonomy')
add_table(d,['Dimension','Levels','Meaning'],[
['Severity','Info / Low / Medium / High / Critical','Potential impact and review urgency; not probability of error. Critical requires policy-defined conditions and human escalation.'],['Confidence','Low / Moderate / High','Strength of evidence for the stated narrow observation; never certainty that a real-world error occurred.'],['Claim state','Observed / Inferred / Reconciled / Unresolved / Verified-by-authority','Observed = direct data fact; inferred = explanation hypothesis; verified-by-authority only when cited authority explicitly establishes it.'],['Disposition','Open / Needs source clarification / Resolved / Not detectable / False-positive after review','A resolution records reviewer, rationale, timestamp and evidence.']])
para(d,'Safe wording examples: “The two published aggregates are not comparable until reporting period is confirmed.” “A potential duplicate key pattern requires source-owner review.” Avoid: “The government data is fake,” “fraud detected,” or “this beneficiary is ineligible.”', 'BDD Callout')

heading(d,'15. AI, RAG and Model Strategy')
para(d,'Model selection is an experiment, not an assumption. Candidate open-weight models may include Qwen2.5 Instruct, Llama 3.1/3.2 Instruct where licence and hardware permit, Mistral Small, Gemma 2/3, and Indic-focused candidates such as Sarvam/OpenHathi-family models where legally available. Confirm current licences, hardware requirements and benchmark suitability before use. The selected model is pinned by repository/commit or checksum, quantization, prompt version and inference settings.')
add_table(d,['Benchmark task','Measure','Gate'],[
['Schema mapping','Top-1 accuracy, top-k recall, abstention precision, reviewer time','Must support abstention; no auto-approval based only on LLM confidence'],['Definition extraction','Span-level support rate, factuality review, ambiguity capture','Every definition must include a source span or be marked proposed'],['Evidence QA','Citation precision, citation completeness, answer faithfulness, unsafe-claim rate','Answers without supporting evidence are refused or marked unsupported'],['Investigation planning','Human-rated usefulness, tool-call validity, unnecessary-call rate','Agent cannot access hidden benchmark content or external-write tools']])
para(d,'RAG pipeline: ingest document → extract structure and page/row locators → chunk by semantic section/table with stable IDs → embed → retrieve with metadata filters (source, date, geography, programme) → optional rerank → provide only bounded evidence to the model → verify each answer citation maps to a retrieved chunk. Qdrant is the V1 vector store; use a local embedding model selected through the same benchmark process.')
heading(d,'15.1 LangGraph orchestration')
codeblock(d,'START → scope_guard → retrieve_evidence → assess_sufficiency\n  → [insufficient: ask_clarification/refuse] | [sufficient: plan]\n  → deterministic_tools (profile/query/compare/lineage) → synthesize_cited_memo\n  → claim_safety_check → human_review_gate (for high/critical) → END')

heading(d,'16. Dashboard and Ask Detective')
add_table(d,['Surface','Required elements'],[
['Forensic dashboard','Source health, data fitness, findings by severity/state, comparison matrix, drift timeline, review queue, filters and no-colour-only status cues'],['Case file','Question, scope, finding timeline, evidence cards, graph path, competing explanations, reproducibility recipe, reviewer decision'],['Dataset page','Manifest, profile, schema mapping, definitions, coverage, caveats, RAG readiness and version comparison'],['Ask Detective','Scoped question box, evidence scope selector, answer with source chips/page-row locators, confidence, assumptions and “investigate next” actions'],['Benchmark console (restricted)','Blinded case registry, execution logs, label reveal control, adjudication and metrics export']])

heading(d,'17. Responsible AI, Security and Provenance')
bullet(d,'Human confirmation is mandatory before a high/critical finding is presented as externally shareable.')
bullet(d,'Separate benchmark labels from runtime storage and prohibit them in retrieval indexes, prompts, logs and UI roles used for investigation.')
bullet(d,'Perform PII/sensitive-data detection on ingest. V1 excludes sensitive raw records unless a documented lawful basis, minimization plan and access control are in place.')
bullet(d,'Store source terms/licence metadata and obey robots/API terms; retain only permitted copies. Cite official source URLs, retrieval date and immutable hash.')
bullet(d,'Use least-privilege service accounts, secret manager/env injection, dependency scanning, TLS in deployed environments, audit logs and role-based review gates.')
bullet(d,'Citation rule: an AI answer must cite each material factual statement to a specific artifact and locator. A citation to a search result, uncaptured web page or another model is insufficient.')

heading(d,'18. Testing, Evaluation and Observability')
heading(d,'18.1 Testing strategy')
add_table(d,['Test class','Examples'],[
['Unit','Date/fiscal-year parsing; unit conversion; hashes; score arithmetic; stable graph-edge construction'],['Contract','Pydantic/OpenAPI request-response validation; engine envelope compatibility'],['Data quality','Great Expectations suites for required columns, ranges, key uniqueness and allowed values'],['Integration','Raw artifact → canonical view → rule run → evidence graph → case export'],['Golden/reproducibility','Pinned fixture hash yields stable profile/findings under pinned rules'],['Adversarial','Prompt injection in source text; copied sources; conflicting dates; null-heavy datasets; duplicate headers; schema rename; manipulated CSV formula cells'],['Human UX','Evidence readability, keyboard navigation, clarity of uncertainty and review workflow']])
heading(d,'18.2 Ground-truth evaluation')
para(d,'Report both case-level and finding-level performance. Define a true positive as a BDD finding that matches an adjudicated benchmark condition under a pre-registered matching rule; define false positives as unsubstantiated findings after reviewer adjudication. Use precision = TP/(TP+FP), recall = TP/(TP+FN), and F1 = 2PR/(P+R). Also report coverage (cases with accessible data), detectability rate, time-to-investigation, evidence completeness, calibration by confidence level and a confusion matrix. Do not report a metric where the denominator or matching policy is undefined.')
heading(d,'18.3 Observability')
para(d,'Capture structured events for ingestion, parser warnings, rule executions, lineage write counts, retrieval IDs/scores, agent tool calls, citation-validation failures, latency, errors and review outcomes. Use OpenTelemetry traces keyed by run_id/case_id; ensure raw sensitive values and hidden labels are redacted from logs.')

heading(d,'19. Deployment, Scalability and Operations')
add_table(d,['Stage','Deployment posture'],[
['Local / hackathon','Docker Compose: web, API, worker, Postgres/PostGIS, Qdrant, object-store emulator and optional local LLM runtime. Seed only approved public fixtures.'],['Team demo','Container registry, managed Postgres/object storage, SSO/RBAC, background queue, scheduled source checks and centralized tracing.'],['Scale-out','Partition raw objects by source/version; queue long profiling jobs; use DuckDB/Polars for batch; cache profiles/embeddings; graduate graph traversal to Neo4j when proven necessary.']])

heading(d,'20. Phased Implementation Roadmap and Definition of Done')
add_table(d,['Phase','Build scope','Definition of done'],[
['Sprint 0 — corpus & governance','Choose one programme; register CAG observations; locate/source-map public artifacts; establish licences and blind-label process.','2–3 candidate cases have manifests, hashes, CAG citations, access status and an approved benchmark protocol.'],
['Sprint 1 — ingestion/profile','Artifact store, manifests, CSV/XLSX/JSON ingestion, profiler, dashboard dataset page.','A raw file can be ingested and profiled reproducibly; profile is linked to source and hash.'],
['Sprint 2 — semantic foundations','Canonical concepts, schema mapping review, geo/temporal/unit normalization.','At least one dataset pair is mapped with evidence and comparability status.'],
['Sprint 3 — forensic engines','Rules, anomaly engine, cross-source comparison, drift and initial graph.','At least one case creates a reproducible finding with a backward lineage path.'],
['Sprint 4 — evidence & AI','RAG evidence store, constrained LangGraph agent, Ask Detective, safe wording checks.','Every material AI sentence is cited or refused; no benchmark-label leakage test passes.'],
['Sprint 5 — evaluation/demo','Blind runs, adjudication, metric report, polish, demo fixtures and video.','Case outcomes are documented; no fabricated performance claims; demo completes from a clean environment.']])

heading(d,'21. Repository Structure and Development Workflow')
codeblock(d,'bdd/\n  apps/web/                  # Next.js UI\n  services/api/              # FastAPI routes and typed contracts\n  packages/contracts/        # OpenAPI/JSON schemas shared across services\n  pipelines/ingestion/       # connectors, parsers, manifests\n  pipelines/forensics/       # profiles, normalization, rules, comparisons\n  pipelines/evaluation/      # blind benchmark runner; restricted fixtures\n  graph/                     # NetworkX adapter and graph schemas\n  rag/                       # chunking, indexing, retrieval, citations\n  data/fixtures/public/      # small approved fixtures only\n  docs/adr/                  # architecture decisions\n  infra/docker/              # compose, migrations, observability\n  tests/                     # unit, integration, golden, adversarial\n  scripts/                   # developer commands\n  .github/workflows/         # CI')
add_table(d,['Practice','Rule'],[
['Branching','Short-lived feature branches; protected main; conventional commits.'],['PR review','Require tests, migration review, security/provenance checklist and screenshot for UI changes.'],['Data changes','No raw source replacement: add a versioned artifact and supersession link.'],['Rules/prompts/models','Version them alongside code; record hashes/config in every run.'],['ADRs','Record decisions on source policy, scoring weights, ontology changes and model choice.']])

heading(d,'22. Sample Forensic Cases')
para(d,'The cases below are illustrative workflow examples unless the team has independently verified the cited audit observation and source links during Sprint 0. Do not represent them as reproduced findings before blinded evaluation.')
add_table(d,['Illustrative case','BDD workflow','Expected safe output'],[
['PM-KISAN audit-linked data-quality case (candidate)','Map CAG observation to an eligible public release; profile identifier/name/bank or status fields only where publicly lawful; compare period/version and run validation rules.','“Audit-linked candidate case. Potential data-quality pattern detected/not detected; public data observability: [state]. Verification required.”'],
['Definition drift across a scheme dashboard release','Compare data dictionary/release note across two publication dates; detect changed denominator or coverage.','“Trend comparison is not directly comparable after [date] because the operational definition appears to have changed.”'],
['Copied aggregate across catalogue and third-party mirror','Trace URL citation/release timestamps and textual/numeric fingerprints.','“Three published claims appear to share a common origin; source count should not be interpreted as independent confirmation.”']])

heading(d,'23. Demo Flow and Competition Positioning')
add_table(d,['Demo beat','What the audience sees'],[
['1. Why it matters','A real audit-backed motivation, framed carefully as historical evidence.'],['2. Ingest','An official public dataset becomes an immutable, cited artifact with hash.'],['3. Understand','Profiler and semantic cards reveal structure, missingness and definitions.'],['4. Investigate','A cross-source/temporal mismatch is aligned before comparison, then flagged with alternatives.'],['5. Prove lineage','Click from finding → rule/query → normalized aggregate → raw row/source.'],['6. Explain responsibly','Ask Detective answers with citations, uncertainty and next verification action.'],['7. Evaluate honestly','Show blinded case protocol and current measured status—not invented precision.']])
para(d,'Competition message: BDD is not “an LLM that finds bad data.” It is a trust infrastructure layer combining deterministic analytics with AI-assisted evidence synthesis. Its signature innovations—semantic drift detection, evidence lineage, false consensus detection and cross-source AI forensics—address the failure modes that make ordinary dashboards and naive RAG unreliable.')

heading(d,'24. Risks and Mitigations')
add_table(d,['Risk','Mitigation'],[
['Public data cannot reproduce a CAG observation','Mark not detectable; document access gap; select a different case without changing labels retroactively.'],['Semantic mismatch produces false contradiction','Require comparability gate; show alternative explanations; reviewer approval for serious findings.'],['LLM hallucination or prompt injection','Evidence-bounded tools, citation validator, source-content isolation, safe claim checker and human gate.'],['Portal/link volatility','Immutable local raw artifacts, manifests, hashes, retrieval timestamps and update monitoring.'],['Scope overload','Protect V1: one domain, batch workflow, 2–3 cases, NetworkX before Neo4j.'],['Sensitive-data exposure','Public-only baseline, minimization, PII detection, redaction, RBAC and no raw values in logs.'],['Evaluation leakage','Separate benchmark store/roles, deny benchmark retrieval, run leakage tests in CI.']])

heading(d,'25. Future Expansion')
para(d,'After V1 demonstrates reproducible value, BDD can become middleware for downstream analytics and RAG: a dataset gateway can publish Data Fitness and RAG Readiness manifests, expose provenance-aware retrieval, block low-fitness fields, attach caveats to answers, and notify consumers when a source version, definition or geography changes. Further expansions include multilingual ontology curation, publisher feedback packets, privacy-preserving secure data rooms, domain-specific validation packs and API integrations for public-data catalogs.')

heading(d,'26. Final Deliverables')
add_table(d,['Deliverable','Minimum contents'],[
['Working prototype','Containerized web/API/worker stack with approved public demo corpus.'],['Corpus & provenance pack','Source registry, manifests, hashes, terms/licence notes, CAG citation registry and benchmark policy.'],['Forensic case files','2–3 blinded-run case bundles with evidence paths and human adjudication.'],['Evaluation report','Methods, matching policy, coverage, case outcomes, metrics where measurable and limitations.'],['Technical documentation','Architecture, API contract, setup guide, data schema, ADRs, operations and security notes.'],['Competition package','Demo script, screenshots/video, one-page summary and responsible-AI statement.']])

heading(d,'Appendix A. Source and Citation Register Template')
add_table(d,['Field','Required entry'],[['source_id','Stable internal ID'],['official_title','Exact source/report/dataset title'],['publisher','Named issuing institution'],['url','Direct official landing page or artifact URL'],['publication/retrieval dates','Both, when available'],['locator','Report page/paragraph, dataset resource/table/row or API params'],['hash','SHA-256 of captured artifact'],['use in BDD','Primary source / alternate / benchmark / context'],['access/licence','Terms or access note'],['verification status','Unverified candidate / source-verified / benchmark-adjudicated']])
heading(d,'Appendix B. Suggested Official Starting Points')
bullet(d,'CAG India reports and report index: https://cag.gov.in/')
bullet(d,'Open Government Data Platform India: https://www.data.gov.in/')
bullet(d,'AIKosh: use its current official catalogue and applicable access/licence terms.')
para(d,'These are discovery starting points, not proof of any specific case. During Sprint 0, capture the exact source URL, page/paragraph or resource locator and artifact hash in the citation register.')

d.core_properties.title='Bharat Data Detective (BDD) — Project Specification'; d.core_properties.subject='AI forensic and evidence trust layer for Indian public data'; d.core_properties.author='BDD Project Team'; d.core_properties.keywords='public data, provenance, forensic AI, RAG, India, CAG'
d.save(OUT)
print(OUT)
